"""Acceptance tests for LIFE OS v0.1. Run from lifeos/: python -m pytest backend/tests -q"""
import copy
import hashlib
import json
import math
import os
import re
import sys

import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from lifeos import calendar as calm  # noqa: E402
from lifeos import metrics  # noqa: E402
from lifeos.catalogs import CatalogError, Catalogs  # noqa: E402
from lifeos.decision import Cand, score  # noqa: E402
from lifeos.engine import EditError, Simulation  # noqa: E402
from lifeos.experiments import analyze, run_one, validate_spec  # noqa: E402
from lifeos.rng import softmax_pick  # noqa: E402
from lifeos.rules import RulesError, schema  # noqa: E402
from lifeos.storage import Store, world_from_bytes, world_to_bytes, world_to_json  # noqa: E402
from lifeos.worldgen import generate  # noqa: E402


def small(seed=99, students=14, featured=False, tracing=True):
    w = generate(seed, students, featured)
    w.tracing = tracing
    return Simulation(w)


def state_hash(w) -> str:
    h = hashlib.sha1()
    for c in sorted(w.chars.values(), key=lambda c: c.id):
        h.update(json.dumps([c.id, c.role, c.grade, {k: round(v, 6) for k, v in c.needs.items()}, round(c.mood, 6),
                             {k: round(v, 6) for k, v in c.skills.items()}, sorted(c.memberships), c.flags.get("partner"),
                             [g["label"] + g["status"] for g in c.goals]]).encode())
    for k in sorted(w.rels):
        r = w.rels[k]
        h.update(json.dumps([k, round(r.affection, 6), round(r.familiarity, 6), round(r.attraction, 6), r.label]).encode())
    return h.hexdigest()


def event_seq(w):
    return [(e["day"], e["type"], e.get("outcome"), tuple(e["actors"]), tuple(e["targets"])) for e in w.history.events.values()]


# --- identity ------------------------------------------------------------------------

def test_characters_have_stable_identities():
    a, b = generate(5, 20, True), generate(5, 20, True)
    assert list(a.chars) == list(b.chars)
    assert [c.name for c in a.chars.values()] == [c.name for c in b.chars.values()]
    sim = Simulation(a)
    ids = set(a.chars)
    names = {cid: c.name for cid, c in a.chars.items()}
    sim.advance_days(10)
    assert ids <= set(a.chars)
    assert all(a.chars[cid].name == names[cid] for cid in ids)
    w2 = world_from_bytes(world_to_bytes(a))
    assert list(w2.chars) == list(a.chars)
    assert len(set(a.chars)) == len(a.chars)


# --- time ---------------------------------------------------------------------------

def test_time_advances_consistently():
    sim = small()
    w = sim.w
    assert (w.day, w.phase) == (0, 0)
    cal = sim.cat.world["calendar"]
    assert sim.phases() == calm.SCHOOL_PHASES  # Sept 1, 2025 is a Monday school day
    sim.advance_days(1)
    assert (w.day, w.phase) == (1, 0)
    sim.advance_days(6)
    assert w.day == 7 and w.phase == 0
    assert calm.weekday(5) == 5 and calm.phases_for(5, cal) == calm.FREE_PHASES
    # stepping walks every phase of the day in order
    seen = []
    d = w.day
    while w.day == d:
        seen.append(sim.current_phase())
        sim.step()
    assert seen == sim.phases() if calm.is_school_day(w.day, cal) else True
    assert seen == calm.phases_for(d, cal)


def test_school_year_rolls_over_and_seniors_graduate():
    sim = small(students=16, tracing=False)
    w = sim.w
    seniors = [c.id for c in w.students() if c.grade == 12]
    juniors = [c.id for c in w.students() if c.grade == 11]
    grad = next(i for i in w.calendar if i["kind"] == "graduation")
    w.day = grad["day"]
    sim._refresh_day()
    sim.advance_days(1)
    assert all(w.chars[s].role == "alumni" for s in seniors)
    assert any(e for e in sim.w.meta["tally"] if e.startswith("graduation."))
    before = len(w.chars)
    w.day = calm.md_to_day(0, (8, 30))  # Aug 30 at the end of the first school year
    sim._refresh_day()
    sim.advance_days(3)
    assert w.year == 1
    assert all(w.chars[j].grade == 12 for j in juniors)
    assert len(w.chars) == before + len(seniors)  # freshmen replace the graduating class
    assert any(i["key"].endswith(":1") for i in w.calendar)  # next year's calendar exists


# --- autonomy -------------------------------------------------------------------------

def test_npcs_act_and_initiate_interactions_without_a_player():
    sim = small()
    sim.advance_days(5)
    w = sim.w
    initiated = [e for e in w.history.events.values() if e["type"].startswith("interact.") and e["actors"]
                 and w.chars[e["actors"][0]].role == "student" and e["targets"]]
    assert len(initiated) >= 10
    assert len({e["actors"][0] for e in initiated}) >= 5
    deciders = {t["char"] for t in w.history.traces.values()}
    assert deciders >= {c.id for c in w.students()}
    assert not any(t.get("player_choice") for t in w.history.traces.values())


def test_relationships_form_and_change_and_are_directed():
    sim = small(students=18)
    w = sim.w
    before = {k: (r.affection, r.familiarity) for k, r in w.rels.items()}
    sim.advance_days(30)
    changed = sum(1 for k, r in w.rels.items() if k in before and abs(r.affection - before[k][0]) > 3)
    new = sum(1 for k in w.rels if k not in before)
    assert changed > 30
    assert new > 0 or changed > 60
    assert any(e["type"] == "milestone" for e in w.history.events.values())
    asym = [k for k, r in w.rels.items() if (k[1], k[0]) in w.rels and abs(r.affection - w.rels[(k[1], k[0])].affection) > 15]
    assert asym, "two people should be able to feel differently about the same relationship"
    labels = {r.label for r in w.rels.values()}
    assert {"acquaintance", "friend"} <= labels


# --- parameters change behavior -----------------------------------------------------------

def _count_actions(rules, days=14, seed=99):
    sim = small(seed=seed, students=16, tracing=False)
    sim.set_rules(rules)
    sim.advance_days(days)
    return sim.w.meta["actions"], sim


def test_decisions_respond_to_parameters():
    sim = small()
    w = sim.w
    c = next(c for c in w.students())
    party = Cand(sim.cat.actions["throw_party"], None, None, "test")
    ctx = sim.char_ctx(c)
    lo = score(sim, c, party, ctx, None)
    sim.set_rules({"motive.approval": 3.0})
    hi = score(sim, c, party, sim.char_ctx(c), None)
    expected = sim.cat.actions["throw_party"]["motives"]["approval"] * c.motives["approval"] / 100 * 15 * 2.0
    assert hi - lo == pytest.approx(expected, rel=1e-6)
    # behavioral: more status-seeking with approval x4 than with approval off
    a_lo, _ = _count_actions({"motive.approval": 0.0})
    a_hi, _ = _count_actions({"motive.approval": 4.0})
    status = ("tease", "gossip", "throw_party", "go_to_party", "spread_rumor", "introduce")
    assert sum(a_hi.get(k, 0) for k in status) > sum(a_lo.get(k, 0) for k in status)
    # romance switch removes romantic behavior entirely
    a_off, s_off = _count_actions({"romance.enabled": False}, days=21)
    assert not any(a_off.get(k) for k in ("flirt", "ask_out", "ask_to_dance"))
    assert not any(k.startswith("milestone.dating") for k in s_off.w.meta["tally"])


def test_temperature_controls_how_deterministic_choices_are():
    sim = small()
    sim.advance_days(2)
    w = sim.w

    def top_prob(T):
        sim.set_rules({"decision.temperature": T})
        c = next(c for c in w.students())
        sim.ctx = sim.slot_ctx("after_school")
        d = sim.decide_for(c.id, "after_school", trace=False)
        sim.ctx = None
        return max(cd.prob for cd in d["cands"])
    assert top_prob(1.0) > top_prob(30.0)


# --- events, history, explainability ---------------------------------------------------------

def test_events_record_the_state_changes_that_actually_happened():
    sim = small()
    w = sim.w
    sim.advance_days(3)
    while sim.current_phase() != "after_school":
        sim.step()
    snap = {c.id: dict(c.needs) for c in w.chars.values()}
    rel_snap = {k: r.affection for k, r in w.rels.items()}
    first = w.history.next_event
    sim.step()  # one decision slot
    logged: dict = {}
    for eid in range(first, w.history.next_event):
        e = w.history.events.get(eid)
        if not e or not e.get("changes"):
            continue
        for entity, fld, b, a, why in e["changes"]:
            assert isinstance(fld, str) and why is not None
            if isinstance(b, (int, float)) and isinstance(a, (int, float)):
                logged[(entity, fld)] = logged.get((entity, fld), 0) + (a - b)
    checked = 0
    for c in w.chars.values():
        for need in ("energy", "social", "fun", "stress"):
            actual = c.needs[need] - snap[c.id][need]
            rec = logged.get((c.id, f"needs.{need}"), 0)
            assert abs(actual - rec) < 1.0, (c.id, need, actual, rec)
            checked += abs(actual) > 0.5
    for (a, b), r in w.rels.items():
        if (a, b) in rel_snap:
            actual = r.affection - rel_snap[(a, b)]
            rec = logged.get((f"{a}>{b}", "affection"), 0)
            assert abs(actual - rec) < 1.0
    assert checked > 5


def test_history_is_consistent_with_itself():
    sim = small()
    sim.advance_days(14)
    h = sim.w.history
    ids = list(h.events)
    assert ids == sorted(ids) or len(ids) > 0
    for eid, e in h.events.items():
        for cid in e["actors"] + e["targets"]:
            assert cid in sim.w.chars
            assert eid in h.by_char[cid]
        for c in e["causes"]:
            assert c < eid
        assert e.get("text")
        if e.get("trace") and e["trace"] in h.traces:
            assert h.traces[e["trace"]]["char"] in e["actors"] + e["targets"] + e.get("involved", [])


def test_important_outcomes_have_inspectable_decision_traces():
    sim = generate(1234, 36, True)
    s = Simulation(sim)
    s.advance_days(21)
    h = sim.history
    important = [e for e in h.events.values() if e.get("trace") and e["importance"] >= 20]
    assert important
    for e in important[:30]:
        tr = h.traces[e["trace"]]
        assert tr["event"] == e["id"] or e.get("data", {}).get("fallback")
        probs = [r["prob"] for r in tr["candidates"]]
        assert abs(sum(probs) - 1) < 0.01
        chosen = [r for r in tr["candidates"] if r["chosen"]]
        assert len(chosen) == 1 and "terms" in chosen[0]
        total = sum(t[1] for t in chosen[0]["terms"])
        assert total == pytest.approx(chosen[0]["utility"], abs=0.2), "terms must add up to the utility"
        assert 0 <= tr["roll"] < 1 and tr["temperature"] > 0
        if e.get("calc"):
            assert all("step" in c for c in e["calc"])
    # the recorded roll reproduces the recorded choice
    for tr in list(h.traces.values())[:200]:
        rows = tr["candidates"]
        if tr.get("player_choice"):
            continue
        # rows are sorted by utility; softmax_pick works on any fixed order, so rebuild probabilities
        us = [r["utility"] for r in rows]
        _, probs = softmax_pick(us, tr["temperature"], 0.5)
        for p, r in zip(probs, rows):
            assert abs(p - r["prob"]) < 0.02


# --- randomness and reproducibility -------------------------------------------------------------

def test_same_seed_reproduces_and_different_seed_diverges():
    a, b, c = small(seed=7), small(seed=7), small(seed=7)
    c.w.seed = 4242
    c = Simulation(c.w)
    for s in (a, b, c):
        s.advance_days(12)
    assert state_hash(a.w) == state_hash(b.w)
    assert event_seq(a.w) == event_seq(b.w)
    assert state_hash(a.w) != state_hash(c.w)


def test_tracing_never_changes_outcomes():
    a, b = small(seed=8, tracing=True), small(seed=8, tracing=False)
    a.advance_days(15)
    b.advance_days(15)
    assert state_hash(a.w) == state_hash(b.w)


def test_save_and_load_preserve_the_world_and_its_future(tmp_path):
    store = Store(str(tmp_path / "t.db"))
    ref = small(seed=11)
    ref.advance_days(16)
    s = small(seed=11)
    s.advance_days(8)
    sid = store.save_world(s.w, "midpoint")
    loaded = Simulation(store.load_world(sid))
    assert state_hash(loaded.w) == state_hash(s.w)
    assert len(loaded.w.history.events) == len(s.w.history.events)
    assert loaded.w.rules == s.w.rules
    loaded.advance_days(8)
    assert state_hash(loaded.w) == state_hash(ref.w), "a loaded save must continue exactly like the original"
    assert [x["name"] for x in store.list_saves()] == ["midpoint"]


# --- configuration validation --------------------------------------------------------------------

def test_invalid_configuration_is_rejected():
    sch = schema()
    for bad in ({"nope.key": 1}, {"motive.approval": 99}, {"motive.approval": "lots"}, {"decision.shortlist_size": 2.5},
                {"romance.enabled": "maybe"}, {"decision.temperature": float("nan")}):
        with pytest.raises(RulesError):
            sch.validate(bad)
    with pytest.raises(RulesError):
        sch.validate_full({"rel.friend_threshold": 70, "rel.close_threshold": 60})
    sim = small()
    before = dict(sim.w.rules)
    with pytest.raises(RulesError):
        sim.set_rules({"motive.approval": 2.0, "decision.temperature": -1})
    assert sim.w.rules == before, "a rejected change must not partially apply"
    with pytest.raises(CatalogError):
        sim.set_action_override("study", {"base": "high"})
    with pytest.raises(CatalogError):
        sim.set_action_override("study", {"traits": {"telepathy": 1}})
    with pytest.raises(EditError):
        sim.edit_character(next(iter(sim.w.chars)), {"traits": {"discipline": 140}})


def test_every_rule_parameter_is_read_by_the_engine():
    src = ""
    root = os.path.join(os.path.dirname(__file__), "..", "lifeos")
    for fn in os.listdir(root):
        if fn.endswith(".py") and fn not in ("api.py", "spreadsheet.py"):
            src += open(os.path.join(root, fn), encoding="utf-8").read()
    missing = []
    for key in schema().params:
        if f'"{key}"' in src:
            continue
        sec, name = key.split(".", 1)
        if sec == "motive" and 'f"motive.{m}"' in src:
            continue
        missing.append(key)
    assert not missing, f"parameters nothing reads: {missing}"


# --- experiments ------------------------------------------------------------------------------------

def test_experiment_runs_are_isolated_and_reproducible():
    base = generate(3, 12, False)
    Simulation(base)
    base_json = world_to_json(base)
    snapshot = copy.deepcopy(base_json)
    job = {"variant": "Baseline", "rules": {}, "seed": 5, "days": 14, "base": base_json}
    v = dict(job, variant="Approval", rules={"motive.approval": 3.0})
    r1 = run_one(job)
    rv = run_one(v)
    r2 = run_one(job)
    assert r1["metrics"] == r2["metrics"] and r1["digests"] == r2["digests"]
    assert base_json == snapshot, "runs must not touch the base scenario"
    assert r1["digests"] != rv["digests"]
    spec = validate_spec({"name": "t", "days": 14, "seeds": [5],
                          "variants": [{"name": "Baseline", "rules": {}}, {"name": "Approval", "rules": {"motive.approval": 3.0}}]})
    res = analyze(spec, [r1, rv])
    assert res["mode"] == "deterministic"
    assert res["divergence"]["Approval"]["5"] is not None
    assert res["characters"] and "diff" in res["characters"][0]
    spec2 = validate_spec(dict(spec, seeds=[5, 6]))
    runs = [r1, rv, run_one(dict(job, seed=6)), run_one(dict(v, seed=6))]
    res2 = analyze(spec2, runs)
    assert res2["mode"] == "repeated"
    m = res2["comparisons"][0]["metrics"]["mutual_friendships"]
    assert m["n"] == 2 and m["ci95"] is not None and "of 2 higher" in m["direction"]
    with pytest.raises(ValueError):
        validate_spec({"days": 3, "seeds": [1], "variants": [{"name": "a", "rules": {}}]})
    with pytest.raises(ValueError):
        validate_spec({"days": 30, "seeds": [1], "variants": [{"name": "a"}, {"name": "b", "rules": {"motive.approval": 50}}]})


# --- player mode, lab edits, workbook -----------------------------------------------------------------

def test_player_choices_are_resolved_by_the_simulation():
    sim = small()
    cid = next(c.id for c in sim.w.students())
    slot, dec = sim.options_for(cid)
    assert slot in calm.DECISION_SLOTS and dec["cands"]
    pick = min(range(len(dec["cands"])), key=lambda i: dec["cands"][i].prob)  # the least likely option
    chosen = dec["cands"][pick].a["id"]
    evs = sim.act_as(cid, pick)
    tr = [t for t in sim.w.history.traces.values() if t["char"] == cid][-1]
    assert tr["player_choice"] == chosen
    mine = [e for e in evs if cid in e["actors"]]
    assert mine or any(cid in e.get("actors", []) + e.get("targets", []) for e in evs) or chosen in ("sit_with", "attend_practice", "attend_event")


def test_lab_edits_are_validated_and_recorded():
    sim = small()
    cid = next(c.id for c in sim.w.students())
    ev = sim.edit_character(cid, {"traits": {"discipline": 90}, "bio": "Edited in the lab."})
    assert sim.w.chars[cid].traits["discipline"] == 90
    assert ev["type"] == "lab" and any(ch[1] == "traits.discipline" for ch in ev["changes"])
    c = sim.create_character({"grade": 10, "archetype": "artist", "first": "Juno"})
    assert c.id in sim.w.chars and any(k[0] == c.id for k in sim.w.rels)


def test_workbook_round_trip():
    from openpyxl import load_workbook
    import io
    from lifeos.spreadsheet import export_workbook, import_workbook
    sim = small()
    data = export_workbook(sim)
    wb = load_workbook(io.BytesIO(data))
    ws = wb["Rules"]
    for row in ws.iter_rows(min_row=2):
        if row[0].value == "motive.approval":
            row[3].value = 2.5
        if row[0].value == "decision.temperature":
            row[3].value = -5
    buf = io.BytesIO()
    wb.save(buf)
    rep = import_workbook(sim, buf.getvalue())
    assert sim.w.rules["motive.approval"] == 2.5
    assert any("decision.temperature" in e for e in rep["errors"])


# --- demonstration scenario -------------------------------------------------------------------------

def test_demo_world_produces_different_lives():
    w = generate(1234, 36, True)
    w.tracing = False
    sim = Simulation(w)
    sim.advance_days(90)
    devon, marcus = w.chars["devon-park"], w.chars["marcus-bell"]
    assert devon.counters.get("practice_skill", 0) + devon.counters.get("study", 0) > \
        marcus.counters.get("practice_skill", 0) + marcus.counters.get("study", 0), "discipline should change how people spend time"
    assert devon.aptitudes["athletic"] < marcus.aptitudes["athletic"]
    m = metrics.compute(w)
    assert m["mutual_friendships"] > 10
    assert m["friendships_formed"] > 10
    t = w.meta["tally"]
    assert sum(v for k, v in t.items() if k.startswith(("tryout.made", "audition."))) > 0
    assert len({tuple(sorted(c.memberships)) for c in w.students()}) > 5
